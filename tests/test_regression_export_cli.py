"""Portable regression tests for TDMS Tool Suite (read -> DataFrame, xlsx export, PNG, CLI parsing).

Run from the project root:

    python -m unittest discover -s tests -v

* Uses only the standard library ``unittest`` plus packages the tool already needs.
* Two small fixtures live in ``tests/fixtures`` (committed). The 4.99 MB capture is optional: set the
  environment variable ``TDMS_LARGE_SAMPLE`` to the ``... 26 09 22 23 09.tdms`` file, or put it in
  ``<project>/test file/``; otherwise those tests are skipped with a message.
* All outputs go to temporary directories. Qt runs with ``QT_QPA_PLATFORM=offscreen``.

The "expected" values below were captured from the implementation at the time the tests were written
(after the pre-optimization baseline commit) and are intentionally pinned: any change in the DataFrame
contents, the xlsx cell contents or the CLI behaviour must be a conscious decision.
"""
import contextlib
import datetime
import io
import os
import re
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("MPLBACKEND", "Agg")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

import tdms

FIXTURES = ROOT / "tests" / "fixtures"
SMALL_08 = FIXTURES / "J3100002102309 - NI-sbRIO-9651-0258f971 - 26 09 22 23 08.tdms"
SMALL_1459 = FIXTURES / "J3100002102309 - NI-sbRIO-9651-0258f971 - 26 09 22 14 59.tdms"
LARGE_SIZE = 4996247
LARGE_SKIP_MESSAGE = (
    "Large TDMS sample not available. Set the environment variable TDMS_LARGE_SAMPLE to the file "
    "'... 26 09 22 23 09.tdms' (4,996,247 bytes) or place it in '<project>/test file/' to run this test."
)
BASE_COLUMNS = ["TIME", "M1 T-Case", "M1 T-Evap", "M1 TSD", "M1 Set Temp"]

# ---- pinned digests of every channel (all 41 columns selected) --------------------------------------
# Captured digests (first value, last value, NaN-skipping sum, NaN count) per channel: SMALL_08 (8, 41) normal
SMALL_08_COLUMNS = ['TIME', 'M1 T-Case', 'M1 T-Evap', 'M1 TSD', 'M1 Set Temp', 'Bypass PID Out', 'Bypass Target', 'Bypass Valve steps', 'Current', 'Current(A)', 'Frequency', 'Inverter RPM', 'M1 AC Load W', 'M1 DC Load W', 'M1 Delta Temp', 'M1 Dynamic SV', 'M1 EEV PID Out', 'M1 EEV Steps', 'M1 EEV Target', 'M1 FORCE', 'M1 FTC200 PV', 'M1 Feedback', 'M1 HEATER %', 'M1 Head State', 'M1 PV', 'M1 Power', 'M1 PowerDrive', 'M1 RH', 'M1 Thermal Load', 'M1 Thermal State', 'P-Discharge(PSIG)', 'P-Suction(PSIG)', 'Power', 'SubCooling(F)', 'Superheat(F)', 'T-Ambient(F)', 'T-Compressor(C)', 'T-Discharge(F)', 'T-Liquid(F)', 'T-Suction(F)', 'Voltage']
SMALL_08_DIGEST = {
    'TIME': (0.0, 7.0, 28.0, 0),
    'M1 T-Case': (30.0, 30.0, 240.0, 0),
    'M1 T-Evap': (30.0, 30.0, 240.0, 0),
    'M1 TSD': (None, None, 0.0, 8),
    'M1 Set Temp': (25.0, 25.0, 200.0, 0),
    'Bypass PID Out': (140.0, 1.0, 284.0, 0),
    'Bypass Target': (140.0, 5.0, 175.0, 0),
    'Bypass Valve steps': (2000.0, 1549.0, 14297.0, 0),
    'Current': (0.0, 0.0, 0.0, 0),
    'Current(A)': (0.0, 1.0, 5.0, 0),
    'Frequency': (0.0, 0.0, 0.0, 0),
    'Inverter RPM': (0.0, 0.0, 0.0, 0),
    'M1 AC Load W': (0.0, 0.0, 0.0, 0),
    'M1 DC Load W': (0.0, 0.0, 0.0, 0),
    'M1 Delta Temp': (-0.0, -0.0, 0.0, 0),
    'M1 Dynamic SV': (25.0, 25.0, 200.0, 0),
    'M1 EEV PID Out': (29.0, 32.0, 246.0, 0),
    'M1 EEV Steps': (29.0, 300.0, 2067.0, 0),
    'M1 EEV Target': (50.0, 50.0, 400.0, 0),
    'M1 FORCE': (0.0, 0.0, 0.0, 0),
    'M1 FTC200 PV': (0.0, 0.0, 0.0, 0),
    'M1 Feedback': (1.0, 1.0, 8.0, 0),
    'M1 HEATER %': (0.0, 0.0, 2.0, 0),
    'M1 Head State': (1.0, 1.0, 8.0, 0),
    'M1 PV': (29.7578125, 29.765625, 238.414, 0),
    'M1 Power': (0.0, 0.0, 0.0, 0),
    'M1 PowerDrive': (0.0, 0.0, 0.0, 0),
    'M1 RH': (None, None, 0.0, 8),
    'M1 Thermal Load': (0.0, 0.0, 0.0, 0),
    'M1 Thermal State': (2.0, 2.0, 16.0, 0),
    'P-Discharge(PSIG)': (133.0, 143.0, 1087.0, 0),
    'P-Suction(PSIG)': (135.0, 132.0, 1073.0, 0),
    'Power': (0.0, 0.0, 0.0, 0),
    'SubCooling(F)': (-6.0, -1.0, -37.0, 0),
    'Superheat(F)': (14.0, 15.0, 113.0, 0),
    'T-Ambient(F)': (64.0, 64.0, 512.0, 0),
    'T-Compressor(C)': (None, None, 0.0, 8),
    'T-Discharge(F)': (79.0, 79.0, 632.0, 0),
    'T-Liquid(F)': (70.0, 69.0, 558.0, 0),
    'T-Suction(F)': (78.0, 78.0, 624.0, 0),
    'Voltage': (0.0, 1.0, 4.0, 0),
}
# SMALL_1459 (16, 41) normal
SMALL_1459_COLUMNS = ['TIME', 'M1 T-Case', 'M1 T-Evap', 'M1 TSD', 'M1 Set Temp', 'Bypass PID Out', 'Bypass Target', 'Bypass Valve steps', 'Current', 'Current(A)', 'Frequency', 'Inverter RPM', 'M1 AC Load W', 'M1 DC Load W', 'M1 Delta Temp', 'M1 Dynamic SV', 'M1 EEV PID Out', 'M1 EEV Steps', 'M1 EEV Target', 'M1 FORCE', 'M1 FTC200 PV', 'M1 Feedback', 'M1 HEATER %', 'M1 Head State', 'M1 PV', 'M1 Power', 'M1 PowerDrive', 'M1 RH', 'M1 Thermal Load', 'M1 Thermal State', 'P-Discharge(PSIG)', 'P-Suction(PSIG)', 'Power', 'SubCooling(F)', 'Superheat(F)', 'T-Ambient(F)', 'T-Compressor(C)', 'T-Discharge(F)', 'T-Liquid(F)', 'T-Suction(F)', 'Voltage']
SMALL_1459_DIGEST = {
    'TIME': (0.0, 15.0, 120.0, 0),
    'M1 T-Case': (30.0, 30.0, 480.0, 0),
    'M1 T-Evap': (30.0, 30.0, 480.0, 0),
    'M1 TSD': (None, None, 0.0, 16),
    'M1 Set Temp': (25.0, 25.0, 400.0, 0),
    'Bypass PID Out': (140.0, 4.0, 153.0, 0),
    'Bypass Target': (140.0, 5.0, 215.0, 0),
    'Bypass Valve steps': (473.0, 72.0, 4893.0, 0),
    'Current': (0.0, 5.0, 5.0, 0),
    'Current(A)': (0.0, 3.0, 18.0, 0),
    'Frequency': (0.0, 45.0, 45.0, 0),
    'Inverter RPM': (0.0, 1341.0, 1341.0, 0),
    'M1 AC Load W': (0.0, 0.0, 0.0, 0),
    'M1 DC Load W': (0.0, 0.0, 0.0, 0),
    'M1 Delta Temp': (-0.0, -0.0, 0.0, 0),
    'M1 Dynamic SV': (25.0, 25.0, 400.0, 0),
    'M1 EEV PID Out': (42.0, 47.0, 721.0, 0),
    'M1 EEV Steps': (42.0, 545.0, 7696.0, 0),
    'M1 EEV Target': (50.0, 50.0, 800.0, 0),
    'M1 FORCE': (0.0, 0.0, 0.0, 0),
    'M1 FTC200 PV': (0.0, 0.0, 0.0, 0),
    'M1 Feedback': (1.0, 1.0, 16.0, 0),
    'M1 HEATER %': (0.0, 0.0, 7.0, 0),
    'M1 Head State': (1.0, 1.0, 16.0, 0),
    'M1 PV': (29.71875, 30.1171875, 479.039, 0),
    'M1 Power': (0.0, 0.0, 0.0, 0),
    'M1 PowerDrive': (0.0, 0.0, 0.0, 0),
    'M1 RH': (None, None, 0.0, 16),
    'M1 Thermal Load': (0.0, 0.0, 0.0, 0),
    'M1 Thermal State': (2.0, 2.0, 32.0, 0),
    'P-Discharge(PSIG)': (143.0, 145.0, 2225.0, 0),
    'P-Suction(PSIG)': (107.0, 124.0, 2056.0, 0),
    'Power': (0.0, 212.0, 212.0, 0),
    'SubCooling(F)': (-1.0, -0.0, -52.0, 0),
    'Superheat(F)': (18.0, 8.0, 96.0, 0),
    'T-Ambient(F)': (64.0, 64.0, 1024.0, 0),
    'T-Compressor(C)': (None, None, 0.0, 16),
    'T-Discharge(F)': (74.0, 73.0, 1174.0, 0),
    'T-Liquid(F)': (69.0, 69.0, 1116.0, 0),
    'T-Suction(F)': (69.0, 67.0, 1082.0, 0),
    'Voltage': (0.0, 46.0, 51.0, 0),
}
# LARGE (30439, 41) normal
LARGE_COLUMNS = ['TIME', 'M1 T-Case', 'M1 T-Evap', 'M1 TSD', 'M1 Set Temp', 'Bypass PID Out', 'Bypass Target', 'Bypass Valve steps', 'Current', 'Current(A)', 'Frequency', 'Inverter RPM', 'M1 AC Load W', 'M1 DC Load W', 'M1 Delta Temp', 'M1 Dynamic SV', 'M1 EEV PID Out', 'M1 EEV Steps', 'M1 EEV Target', 'M1 FORCE', 'M1 FTC200 PV', 'M1 Feedback', 'M1 HEATER %', 'M1 Head State', 'M1 PV', 'M1 Power', 'M1 PowerDrive', 'M1 RH', 'M1 Thermal Load', 'M1 Thermal State', 'P-Discharge(PSIG)', 'P-Suction(PSIG)', 'Power', 'SubCooling(F)', 'Superheat(F)', 'T-Ambient(F)', 'T-Compressor(C)', 'T-Discharge(F)', 'T-Liquid(F)', 'T-Suction(F)', 'Voltage']
LARGE_DIGEST = {
    'TIME': (0.0, 30445.0, 463361656.0, 0),
    'M1 T-Case': (30.0, 25.0, 761011.0, 0),
    'M1 T-Evap': (30.0, 24.0, 717849.0, 0),
    'M1 TSD': (None, None, 599.0, 30421),
    'M1 Set Temp': (25.0, 25.0, 760975.0, 0),
    'Bypass PID Out': (131.0, 105.0, 3168105.0, 0),
    'Bypass Target': (140.0, 5.0, 152465.0, 0),
    'Bypass Valve steps': (1490.0, 446.0, 13008365.0, 0),
    'Current': (0.0, 5.0, 151860.0, 0),
    'Current(A)': (0.0, 10.0, 304251.0, 0),
    'Frequency': (0.0, 173.0, 5258386.0, 0),
    'Inverter RPM': (0.0, 5199.0, 158024373.0, 0),
    'M1 AC Load W': (0.0, 0.0, 0.0, 0),
    'M1 DC Load W': (0.0, 0.0, 0.0, 0),
    'M1 Delta Temp': (0.0, 2.0, 44904.0, 0),
    'M1 Dynamic SV': (25.0, 25.0, 760975.0, 0),
    'M1 EEV PID Out': (32.0, 28.0, 848327.0, 0),
    'M1 EEV Steps': (32.0, 252.0, 7771145.0, 0),
    'M1 EEV Target': (50.0, 50.0, 1521950.0, 0),
    'M1 FORCE': (0.0, 0.0, 0.0, 0),
    'M1 FTC200 PV': (0.0, 0.0, 0.0, 0),
    'M1 Feedback': (1.0, 1.0, 30439.0, 0),
    'M1 HEATER %': (0.0, 48.0, 1521183.0, 0),
    'M1 Head State': (1.0, 1.0, 30439.0, 0),
    'M1 PV': (30.0546875, 25.0, 761027.156, 0),
    'M1 Power': (0.0, 0.0, 0.0, 0),
    'M1 PowerDrive': (0.0, 0.0, 0.0, 0),
    'M1 RH': (None, None, 0.0, 30439),
    'M1 Thermal Load': (0.0, 0.0, 0.0, 0),
    'M1 Thermal State': (2.0, 2.0, 60878.0, 0),
    'P-Discharge(PSIG)': (136.0, 152.0, 4665315.0, 0),
    'P-Suction(PSIG)': (138.0, 4.0, 154129.0, 0),
    'Power': (0.0, 750.0, 22981495.0, 0),
    'SubCooling(F)': (-4.0, -0.0, 5552.0, 0),
    'Superheat(F)': (5.0, 9.0, 276403.0, 0),
    'T-Ambient(F)': (64.0, 65.0, 1959049.0, 0),
    'T-Compressor(C)': (None, None, 0.0, 30439),
    'T-Discharge(F)': (80.0, 87.0, 2683523.0, 0),
    'T-Liquid(F)': (70.0, 72.0, 2191125.0, 0),
    'T-Suction(F)': (71.0, -31.0, -898372.0, 0),
    'Voltage': (0.0, 163.0, 4953201.0, 0),
}

_app = None


def setUpModule():
    global _app
    _app = QApplication.instance() or QApplication([])


def tearDownModule():
    global _app
    _app = None


def find_large_sample():
    candidates = []
    env = os.environ.get("TDMS_LARGE_SAMPLE")
    if env:
        candidates.append(Path(env))
    candidates.extend(sorted((ROOT / "test file").glob("*23 09.tdms")))
    for candidate in candidates:
        if candidate.is_file() and candidate.stat().st_size == LARGE_SIZE:
            return candidate
    return None


def make_window():
    return tdms.TDMSFilterApp()


def run_to_completion(window, timeout_ms=60000):
    """process_files is asynchronous (QThread): run the Qt event loop until the window reports it is finished."""
    from PySide6.QtCore import QEventLoop, QTimer
    loop = QEventLoop()
    results = []
    window.processing_finished.connect(lambda summary: (results.append(summary), loop.quit()))
    QTimer.singleShot(timeout_ms, loop.quit)
    with contextlib.redirect_stdout(io.StringIO()):
        window.process_files()
        loop.exec()
    if not results:
        raise AssertionError("process_files did not finish within %d ms" % timeout_ms)
    return results[0]


def quiet(func, *args, **kwargs):
    """Run func while swallowing the tool's progress prints."""
    with contextlib.redirect_stdout(io.StringIO()):
        return func(*args, **kwargs)


def load_frame(window, path, all_channels):
    window.file_paths = [str(path)]
    quiet(window._rebuild_checklist_from_all_files)
    if all_channels:
        window._select_all_channels()
    else:
        window._deselect_all_channels()
    return quiet(window.read_tdms_file, str(path))


def cells_of(xlsx_path):
    """Return (sheetnames, rows) where each cell is (value, data_type, number_format)."""
    workbook = load_workbook(xlsx_path)
    try:
        sheet = workbook.active
        rows = [[(c.value, c.data_type, c.number_format) for c in row] for row in sheet.iter_rows()]
        return workbook.sheetnames, rows
    finally:
        workbook.close()


def expected_cell_for(value):
    """Characterized current behaviour for a float32 DataFrame value (NaN -> empty inline-string cell)."""
    if value != value:
        return (None, "inlineStr", "General")
    return (float(value), "n", "General")


class ReadRegressionMixin:
    def check_digest(self, df, columns, digest):
        self.assertEqual(list(df.columns), columns)
        for name, (first, last, total, nan_count) in digest.items():
            series = df[name]
            self.assertEqual(int(series.isna().sum()), nan_count, name)
            if nan_count == len(series):
                continue
            for actual, expected in ((float(series.iloc[0]), first), (float(series.iloc[-1]), last)):
                if expected is None:  # the digest records NaN as None
                    self.assertTrue(actual != actual, f"{name}: expected NaN, got {actual}")
                else:
                    self.assertAlmostEqual(actual, expected, places=3, msg=name)
            self.assertAlmostEqual(float(np.nansum(series.to_numpy(dtype="float64"))), total, places=2, msg=name)


class ReadRegressionTests(ReadRegressionMixin, unittest.TestCase):
    def test_small_fixtures_base_columns_dtype_and_rows(self):
        for path, rows in ((SMALL_08, 8), (SMALL_1459, 16)):
            with self.subTest(fixture=path.name[-12:]):
                df, status = load_frame(make_window(), path, all_channels=False)
                self.assertEqual(status, "normal")
                self.assertEqual(list(df.columns), BASE_COLUMNS)
                self.assertEqual(df.shape, (rows, 5))
                self.assertTrue(all(str(t) == "float32" for t in df.dtypes))
                self.assertEqual(df["TIME"].tolist(), [float(i) for i in range(rows)])
                self.assertEqual(df["M1 T-Case"].tolist(), [30.0] * rows)
                self.assertEqual(df["M1 Set Temp"].tolist(), [25.0] * rows)
                self.assertEqual(int(df["M1 TSD"].isna().sum()), rows)  # channel exists but has no values

    def test_small_fixtures_all_channels(self):
        for path, rows, columns, digest in (
            (SMALL_08, 8, SMALL_08_COLUMNS, SMALL_08_DIGEST),
            (SMALL_1459, 16, SMALL_1459_COLUMNS, SMALL_1459_DIGEST),
        ):
            with self.subTest(fixture=path.name[-12:]):
                df, status = load_frame(make_window(), path, all_channels=True)
                self.assertEqual(status, "normal")
                self.assertEqual(df.shape, (rows, 41))
                self.assertTrue(all(str(t) == "float32" for t in df.dtypes))
                self.check_digest(df, columns, digest)

    def test_checklist_lists_36_extra_channels_without_base_channels(self):
        window = make_window()
        window.file_paths = [str(SMALL_08)]
        quiet(window._rebuild_checklist_from_all_files)
        names = list(window.channel_checkboxes)
        self.assertEqual(len(names), 36)
        self.assertEqual(names, sorted(names))
        self.assertFalse(set(names) & set(BASE_COLUMNS))

    def test_large_sample_all_channels(self):
        large = find_large_sample()
        if large is None:
            self.skipTest(LARGE_SKIP_MESSAGE)
        df, status = load_frame(make_window(), large, all_channels=True)
        self.assertEqual(status, "normal")
        self.assertEqual(df.shape, (30439, 41))
        self.assertTrue(all(str(t) == "float32" for t in df.dtypes))
        self.check_digest(df, LARGE_COLUMNS, LARGE_DIGEST)


class ExcelExportRegressionTests(unittest.TestCase):
    """Pins what export_to_excel writes today (cell level), so a faster implementation can be proven identical."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def copy_fixture(self, source, subdir=""):
        folder = self.tmp / subdir if subdir else self.tmp
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / source.name
        shutil.copyfile(source, target)
        return target

    def export(self, window, df, source_copy):
        folder = source_copy.parent
        before = set(folder.iterdir())
        window.generated_excel_paths.clear()
        window.export_to_excel(df, str(source_copy))
        created = sorted(set(folder.iterdir()) - before)
        self.assertEqual(len(created), 1, created)
        return created[0]

    def assert_matches_frame(self, xlsx, df):
        names, rows = cells_of(xlsx)
        self.assertEqual(names, ["Sheet1"])
        self.assertEqual(len(rows), len(df) + 1)
        self.assertEqual([c[0] for c in rows[0]], list(df.columns))
        self.assertTrue(all(c[1] == "s" and c[2] == "General" for c in rows[0]))
        for r in range(len(df)):
            self.assertEqual(len(rows[r + 1]), df.shape[1])
            for j in range(df.shape[1]):
                expected = expected_cell_for(df.iat[r, j])
                got = rows[r + 1][j]
                self.assertEqual(got, expected, f"row {r + 2} col {df.columns[j]}")

    def test_filename_location_sheet_and_registration(self):
        source = self.copy_fixture(SMALL_08)
        window = make_window()
        df, _ = load_frame(window, source, all_channels=False)
        out = self.export(window, df, source)
        self.assertEqual(out.parent, source.parent)
        self.assertRegex(out.name, r"^" + re.escape(source.stem) + r"_\d{8}_\d{6}\.xlsx$")
        self.assertEqual(window.generated_excel_paths, [str(out)])
        datetime.datetime.strptime(out.stem[-15:], "%Y%m%d_%H%M%S")

    def test_cells_match_dataframe_base_and_all_channels(self):
        for fixture in (SMALL_08, SMALL_1459):
            for all_channels in (False, True):
                with self.subTest(fixture=fixture.name[-12:], all_channels=all_channels):
                    source = self.copy_fixture(fixture, f"{fixture.name[-12:-5]}_{all_channels}")
                    window = make_window()
                    df, _ = load_frame(window, source, all_channels=all_channels)
                    out = self.export(window, df, source)
                    self.assert_matches_frame(out, df)

    def test_all_empty_columns_keep_header_and_empty_cells(self):
        source = self.copy_fixture(SMALL_08)
        window = make_window()
        df, _ = load_frame(window, source, all_channels=True)
        empty_columns = [c for c in df.columns if df[c].isna().all()]
        self.assertEqual(empty_columns, ["M1 TSD", "M1 RH", "T-Compressor(C)"])
        _, rows = cells_of(self.export(window, df, source))
        for name in empty_columns:
            j = list(df.columns).index(name)
            self.assertEqual(rows[0][j][0], name)
            for r in range(1, len(rows)):
                self.assertEqual(rows[r][j], (None, "inlineStr", "General"))

    def test_mixed_dtype_frame_cell_types(self):
        """Characterizes cell type/format for non-float32 inputs (pandas to_excel semantics)."""
        df = pd.DataFrame({
            "f32": np.array([1.5, np.nan, 0.1], dtype="float32"),
            "f64": [1.1, np.nan, 3.0],
            "int": [1, 2, 3],
            "obj": [None, "x", 1],
            "str": ["a", "b", None],
            "bool": [True, False, True],
            "dt": pd.to_datetime(["2026-01-02 03:04:05", None, "2026-01-03 00:00:00"]),
            "allnone": [None, None, None],
            "inf": [np.inf, -np.inf, 2.0],
        })
        source = self.copy_fixture(SMALL_08)
        window = make_window()
        out = self.export(window, df, source)
        names, rows = cells_of(out)
        self.assertEqual(names, ["Sheet1"])
        self.assertEqual([c[0] for c in rows[0]], list(df.columns))
        empty = (None, "inlineStr", "General")
        dt_fmt = "YYYY-MM-DD HH:MM:SS"
        expected = [
            [(1.5, "n", "General"), (1.1, "n", "General"), (1, "n", "General"), empty, ("a", "s", "General"),
             (True, "b", "General"), (datetime.datetime(2026, 1, 2, 3, 4, 5), "d", dt_fmt), empty,
             ("inf", "s", "General")],
            [empty, empty, (2, "n", "General"), ("x", "s", "General"), ("b", "s", "General"),
             (False, "b", "General"), empty, empty, ("-inf", "s", "General")],
            [(0.1000000014901161, "n", "General"), (3, "n", "General"), (3, "n", "General"), (1, "n", "General"),
             empty, (True, "b", "General"), (datetime.datetime(2026, 1, 3, 0, 0), "d", dt_fmt), empty,
             (2, "n", "General")],
        ]
        for r, expected_row in enumerate(expected, start=1):
            for j, expected_cell in enumerate(expected_row):
                got = rows[r][j]
                if isinstance(expected_cell[0], float):
                    self.assertAlmostEqual(got[0], expected_cell[0], places=12, msg=f"r{r} c{j}")
                    self.assertEqual(got[1:], expected_cell[1:], f"r{r} c{j}")
                else:
                    self.assertEqual(got, expected_cell, f"r{r} c{j}")

    def test_all_float_frame_with_nan_and_inf(self):
        """Frames whose columns are all float (the tool's real data) take the streaming fast path."""
        df = pd.DataFrame({
            "f64": [1.5, np.nan, np.inf, -np.inf, 1e300, 0.1 + 0.2],
            "f32": np.array([0.5, 1.25, np.nan, 3.0, -2.0, 0.1], dtype="float32"),
        })
        source = self.copy_fixture(SMALL_08)
        window = make_window()
        names, rows = cells_of(self.export(window, df, source))
        self.assertEqual(names, ["Sheet1"])
        self.assertEqual([c[0] for c in rows[0]], ["f64", "f32"])
        self.assertTrue(all(c[1] == "s" for c in rows[0]))
        empty = (None, "inlineStr", "General")
        expected = [
            [(1.5, "n", "General"), (0.5, "n", "General")],
            [empty, (1.25, "n", "General")],
            [("inf", "s", "General"), empty],
            [("-inf", "s", "General"), (3.0, "n", "General")],
            [(1e300, "n", "General"), (-2.0, "n", "General")],
            [(0.3, "n", "General"), (0.1000000014901161, "n", "General")],
        ]
        self.assertEqual(len(rows), 7)
        for r, expected_row in enumerate(expected, start=1):
            for j, expected_cell in enumerate(expected_row):
                got = rows[r][j]
                if isinstance(expected_cell[0], float):
                    self.assertAlmostEqual(got[0] / expected_cell[0], 1.0, places=12, msg=f"r{r} c{j}")
                    self.assertEqual(got[1:], expected_cell[1:], f"r{r} c{j}")
                else:
                    self.assertEqual(got, expected_cell, f"r{r} c{j}")

    def test_header_only_frame(self):
        source = self.copy_fixture(SMALL_08)
        window = make_window()
        out = self.export(window, pd.DataFrame({"a": [], "b": []}), source)
        names, rows = cells_of(out)
        self.assertEqual(names, ["Sheet1"])
        self.assertEqual([[c[0] for c in row] for row in rows], [["a", "b"]])

    def test_timezone_aware_datetime_is_rejected(self):
        df = pd.DataFrame({"t": pd.to_datetime(["2026-01-01"]).tz_localize("UTC")})
        source = self.copy_fixture(SMALL_08)
        window = make_window()
        with self.assertRaises(ValueError):
            window.export_to_excel(df, str(source))
        self.assertEqual(window.generated_excel_paths, [])  # failed export is never reported as success

    def test_large_sample_base_columns(self):
        large = find_large_sample()
        if large is None:
            self.skipTest(LARGE_SKIP_MESSAGE)
        source = self.copy_fixture(large)
        window = make_window()
        df, _ = load_frame(window, source, all_channels=False)
        self.assert_matches_frame(self.export(window, df, source), df)


class ProcessFilesIntegrationTests(unittest.TestCase):
    def test_process_files_writes_xlsx_and_png_next_to_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / SMALL_1459.name
            shutil.copyfile(SMALL_1459, source)
            window = make_window()
            window.file_paths = [str(source)]
            window.enable_excel.setChecked(True)
            window.auto_open_excel.setChecked(False)
            window.auto_open_plot.setChecked(False)
            boom = AssertionError("unexpected message box")
            with patch.object(tdms.QMessageBox, "critical", side_effect=boom), \
                    patch.object(tdms.QMessageBox, "information", side_effect=boom), \
                    patch.object(tdms.QMessageBox, "warning", side_effect=boom):
                run_to_completion(window)
            xlsx = sorted(Path(directory).glob("*.xlsx"))
            png = sorted(Path(directory).glob("*.png"))
            self.assertEqual(len(xlsx), 1)
            self.assertEqual(len(png), 1)
            self.assertEqual(window.generated_excel_paths, [str(xlsx[0])])
            self.assertEqual(window.generated_image_paths, [str(png[0])])
            self.assertGreater(png[0].stat().st_size, 0)
            self.assertEqual(png[0].read_bytes()[:8], b"\x89PNG\r\n\x1a\n")
            names, rows = cells_of(xlsx[0])
            self.assertEqual(names, ["Sheet1"])
            self.assertEqual(len(rows), 17)  # header + 16 data rows
            self.assertEqual([c[0] for c in rows[0]], BASE_COLUMNS)

    def test_process_files_without_excel_still_writes_png(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / SMALL_08.name
            shutil.copyfile(SMALL_08, source)
            window = make_window()
            window.file_paths = [str(source)]
            window.enable_excel.setChecked(False)
            window.auto_open_plot.setChecked(False)
            run_to_completion(window)
            self.assertEqual(list(Path(directory).glob("*.xlsx")), [])
            self.assertEqual(len(list(Path(directory).glob("*.png"))), 1)


class CommandLineCharacterizationTests(unittest.TestCase):
    """CHARACTERIZATION of the CURRENT behaviour of command_line_tdms_path.

    These tests freeze today's behaviour (including quirks). They do NOT claim it is the ideal behaviour:
    e.g. literal quotes, flags such as --help and extra arguments are handled the way they are today.
    The FSE Workbench integration calls: tdms.exe <filePath>.
    """

    NOT_TDMS = "啟動參數不是 .tdms 檔案：\n"
    NOT_FOUND = "找不到啟動參數指定的 TDMS 檔案：\n"

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)

    def touch(self, name):
        path = self.dir / name
        path.write_bytes(b"")
        return path

    def test_no_argument(self):
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe"]), (None, None))
        self.assertEqual(tdms.command_line_tdms_path([]), (None, None))

    def test_single_valid_path_is_resolved(self):
        path = self.touch("plain.tdms")
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(path)]), (str(path.resolve()), None))

    def test_path_with_spaces_and_chinese(self):
        path = self.touch("含 空格 與中文 測試 file.tdms")
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(path)]), (str(path.resolve()), None))

    def test_extension_check_is_case_insensitive(self):
        path = self.touch("UPPER.TDMS")
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(path)]), (str(path.resolve()), None))

    def test_only_first_argument_is_used_current_behavior(self):
        first = self.touch("first.tdms")
        second = self.touch("second.tdms")
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(first), str(second)]),
                         (str(first.resolve()), None))
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(first), "--whatever"]),
                         (str(first.resolve()), None))

    def test_wrong_extension(self):
        path = self.touch("note.txt")
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(path)]),
                         (None, f"{self.NOT_TDMS}{path}"))

    def test_missing_file(self):
        path = self.dir / "missing.tdms"
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(path)]),
                         (None, f"{self.NOT_FOUND}{path}"))

    def test_directory_arguments(self):
        folder = self.dir / "folder_without_suffix"
        folder.mkdir()
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(folder)]),
                         (None, f"{self.NOT_TDMS}{folder}"))
        tricky = self.dir / "folder.tdms"
        tricky.mkdir()
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", str(tricky)]),
                         (None, f"{self.NOT_FOUND}{tricky}"))

    def test_literal_quotes_are_not_stripped_current_behavior(self):
        path = self.touch("quoted.tdms")
        quoted = f'"{path}"'
        self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", quoted]),
                         (None, f"{self.NOT_TDMS}{Path(quoted)}"))

    def test_flags_are_treated_as_file_names_current_behavior(self):
        for flag in ("--help", "-h", "/?"):
            with self.subTest(flag=flag):
                self.assertEqual(tdms.command_line_tdms_path(["tdms.exe", flag]),
                                 (None, f"{self.NOT_TDMS}{Path(flag)}"))

    def test_blank_argument(self):
        result = tdms.command_line_tdms_path(["tdms.exe", "  "])
        self.assertIsNone(result[0])
        self.assertTrue(result[1].startswith(self.NOT_TDMS))


if __name__ == "__main__":
    unittest.main()
